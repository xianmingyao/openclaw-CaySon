# -*- coding: utf-8 -*-
"""周报生成器 - 2026年第24周"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from datetime import datetime

def set_run_font(run, font_name='微软雅黑', font_size=11, bold=False, color=None):
    """设置run的字体格式"""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    # 中文字体
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def add_heading(doc, text, level=1, font_size=16, bold=True, color=(0, 0, 0)):
    """添加标题"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, font_size=font_size, bold=bold, color=color)
    return p

def add_para(doc, text, font_size=11, bold=False, indent=False, bullet=False):
    """添加段落"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.3)
    if bullet:
        p.style = 'List Bullet'
    run = p.add_run(text)
    set_run_font(run, font_size=font_size, bold=bold)
    return p

def add_table_row(table, col1, col2, col3, col4="", bold_row=False):
    """添加表格行"""
    row = table.add_row()
    cells = [row.cells[0], row.cells[1], row.cells[2], row.cells[3]]
    texts = [col1, col2, col3, col4]
    for i, (cell, text) in enumerate(zip(cells, texts)):
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.bold = bold_row
                run.font.name = '微软雅黑'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return row

def create_weekly_report():
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # ===== 标题 =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('工 作 周 报')
    set_run_font(run, font_size=18, bold=True, color=(0, 51, 102))

    # 副标题（时间）
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('2026年第24周 (2026.06.08 - 2026.06.12)')
    set_run_font(run, font_size=12, bold=False, color=(102, 102, 102))

    doc.add_paragraph()  # 空行

    # ===== 一、上周工作回顾 =====
    add_heading(doc, '一、上周工作回顾', font_size=14, bold=True, color=(0, 51, 102))

    # 工作内容表格
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    hdr_cells = table.rows[0].cells
    headers = ['日期', '工作模块', '完成进度', '主要工作内容']
    for i, (cell, header) in enumerate(zip(hdr_cells, headers)):
        cell.text = header
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = '微软雅黑'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        # 表头背景色
        from docx.oxml.ns import nsmap
        from docx.oxml import OxmlElement
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '003366')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:val'), 'clear')
        tcPr.append(shd)

    # 数据行
    work_records = [
        ('2026.06.08', '京麦商品发布智能体', '[97%]', '智能体构建与知识库技术备忘录开发；京麦商品批量自动化发布调试；解决WebView表单输入CEF兼容问题'),
        ('2026.06.09', '京麦商品发布智能体', '[98%]', '批量商品上架测试完成；修复断点续传Bug；编写调教AI教程'),
        ('2026.06.10', '京麦商品发布智能体', '[98%]', 'DesktopAgent+UFO+Dispatcher+ConditionalRouting最后一公里集成验证；UfoDesktopBackend质量修复；RowDispatcher多行并发调度器开发'),
        ('2026.06.11', '京麦商品发布智能体', '[98%]', '批量商品上架测试修复；处理付总小龙虾异常；局域网算力串行链接'),
        ('2026.06.12', '京麦商品发布智能体', '[98%]', '批量商品上架Bug修复并继续测试；局域网算力串行链接部署'),
    ]

    for record in work_records:
        row = table.add_row()
        for i, (cell, text) in enumerate(zip(row.cells, record)):
            cell.text = text
            for p in cell.paragraphs:
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.name = '微软雅黑'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    doc.add_paragraph()

    # ===== 关键技术问题与解决 =====
    add_heading(doc, '二、关键技术问题与解决方案', font_size=14, bold=True, color=(0, 51, 102))

    issues = [
        ('京麦WebView表单输入问题', '京麦内部使用WebView（CEF），其表单输入框对普通鼠标点击和键盘输入不响应', '使用Windows UIA、Win32、WinCOM原生控件做替代兜底'),
        ('断点续传Bug', '批量商品发布中断后无法从断点继续', '修复断点续传逻辑'),
        ('多行并发调度', '批量商品发布需要支持多行并发', '开发RowDispatcher多行并发调度器'),
    ]

    for title, problem, solution in issues:
        p = doc.add_paragraph()
        run = p.add_run(f'• {title}：')
        set_run_font(run, font_size=11, bold=True)
        run = p.add_run(f'问题：{problem} → 解决：{solution}')
        set_run_font(run, font_size=11)

    doc.add_paragraph()

    # ===== 三、本周工作计划 =====
    add_heading(doc, '三、本周工作计划', font_size=14, bold=True, color=(0, 51, 102))

    plans = [
        ('1', '京麦商品发布智能体的大模型提示词评估决策问题修复', '持续优化智能体的提示词工程，提升大模型在商品发布场景下的决策准确性'),
        ('2', '局域网算力串行链接部署', '在公司内部空闲电脑上安装Linux系统，实现算力串接'),
        ('3', '文生图部署与实现效果摸底', '提前研究文生图技术的部署方案和实际效果评估'),
    ]

    for num, title, desc in plans:
        p = doc.add_paragraph()
        run = p.add_run(f'{num}. {title}')
        set_run_font(run, font_size=11, bold=True)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        run = p.add_run(f'   {desc}')
        set_run_font(run, font_size=11)

    doc.add_paragraph()

    # ===== 四、备注 =====
    add_heading(doc, '四、备注', font_size=14, bold=True, color=(0, 51, 102))

    notes = [
        '京麦商品发布智能体整体完成度[98%]，预计本周完成收尾工作',
        '处理付总小龙虾异常，保持业务稳定运行',
        '算力串行链接将为后续AI推理任务提供更强的计算支持',
    ]

    for note in notes:
        p = doc.add_paragraph()
        run = p.add_run(f'• {note}')
        set_run_font(run, font_size=11)

    doc.add_paragraph()
    doc.add_paragraph()

    # 签名
    signature = doc.add_paragraph()
    signature.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = signature.add_run(f'报告人：CaySon\n日期：{datetime.now().strftime("%Y-%m-%d")}')
    set_run_font(run, font_size=11, color=(102, 102, 102))

    # 保存
    output_path = r'E:\workspace\周报_2026年第24周.docx'
    doc.save(output_path)
    print(f'周报已生成：{output_path}')
    return output_path

if __name__ == '__main__':
    create_weekly_report()
