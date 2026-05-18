"""Excel 导入服务。"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import openpyxl

from jingmai_publish.repositories.upload_job import UploadJobRepository


@dataclass(slots=True)
class ParsedExcelRow:
    """Excel 单行解析结果。"""

    row_no: int
    apply_business: str | None
    product_name: str
    brand: str | None
    model: str | None
    length_mm: Decimal | None
    width_mm: Decimal | None
    height_mm: Decimal | None
    weight_kg: Decimal | None
    unit_name: str | None
    jd_sale_price: Decimal | None
    purchase_price: Decimal | None
    market_price: Decimal | None
    jd_item_url: str
    qualification_pdf_path: str | None
    product_summary: str | None
    remark: str | None
    item_type: str


class ExcelIngestService:
    """负责读取 Excel 并生成批次行数据。"""

    def __init__(self, upload_job_repo: UploadJobRepository) -> None:
        """注入导入仓库。"""

        self.upload_job_repo = upload_job_repo

    @staticmethod
    def _normalize_decimal(value: Any) -> Decimal | None:
        """将 Excel 单元格值归一化为 Decimal。"""

        if value in (None, ""):
            return None
        return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    @staticmethod
    def calculate_purchase_price(jd_sale_price: Decimal) -> Decimal:
        """按业务规则计算采购价。

        规则：采购价 = 京东价 * 0.95
        """

        return (jd_sale_price * Decimal("0.95")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    @staticmethod
    def calculate_market_price(jd_sale_price: Decimal) -> Decimal:
        """按业务规则计算市场价。

        规则：市场价 = 京东价 / 0.85
        """

        return (jd_sale_price / Decimal("0.85")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    def load_headers(self, excel_path: str | Path, sheet_name: str = "上架模板") -> list[str]:
        """读取模板表头。"""

        workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        headers = [value for value in next(sheet.iter_rows(min_row=3, max_row=3, values_only=True)) if value]
        workbook.close()
        return [str(item) for item in headers]

    def parse_rows(self, excel_path: str | Path, sheet_name: str = "上架模板") -> list[ParsedExcelRow]:
        """读取并解析 Excel 数据行。"""

        workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        sheet = workbook[sheet_name]
        parsed_rows: list[ParsedExcelRow] = []

        for raw_row in sheet.iter_rows(min_row=4, values_only=True):
            if raw_row is None or not any(value not in (None, "") for value in raw_row[:15]):
                continue

            jd_sale_price = self._normalize_decimal(raw_row[10])
            purchase_price = self.calculate_purchase_price(jd_sale_price) if jd_sale_price is not None else None
            market_price = self.calculate_market_price(jd_sale_price) if jd_sale_price is not None else None

            parsed_rows.append(
                ParsedExcelRow(
                    row_no=int(raw_row[0]),
                    apply_business=str(raw_row[1]) if raw_row[1] not in (None, "") else None,
                    product_name=str(raw_row[2]),
                    brand=str(raw_row[3]) if raw_row[3] not in (None, "") else None,
                    model=str(raw_row[4]) if raw_row[4] not in (None, "") else None,
                    length_mm=self._normalize_decimal(raw_row[5]),
                    width_mm=self._normalize_decimal(raw_row[6]),
                    height_mm=self._normalize_decimal(raw_row[7]),
                    weight_kg=self._normalize_decimal(raw_row[8]),
                    unit_name=str(raw_row[9]) if raw_row[9] not in (None, "") else None,
                    jd_sale_price=jd_sale_price,
                    purchase_price=purchase_price,
                    market_price=market_price,
                    jd_item_url=str(raw_row[11]),
                    qualification_pdf_path=str(raw_row[12]) if raw_row[12] not in (None, "") else None,
                    product_summary=str(raw_row[13]) if raw_row[13] not in (None, "") else None,
                    remark=str(raw_row[14]) if raw_row[14] not in (None, "") else None,
                    item_type="single",
                )
            )

        workbook.close()
        return parsed_rows

    def ingest_rows(self, job_id: str, rows: list[ParsedExcelRow]) -> list[int]:
        """将解析结果入库为导入行数据。"""

        item_ids: list[int] = []
        for row in rows:
            item = self.upload_job_repo.create_job_item(
                job_id=job_id,
                row_no=row.row_no,
                apply_business=row.apply_business,
                product_name=row.product_name,
                brand=row.brand,
                model=row.model,
                length_mm=row.length_mm,
                width_mm=row.width_mm,
                height_mm=row.height_mm,
                weight_kg=row.weight_kg,
                unit_name=row.unit_name,
                jd_sale_price=row.jd_sale_price,
                purchase_price=row.purchase_price,
                market_price=row.market_price,
                jd_item_url=row.jd_item_url,
                qualification_pdf_path=row.qualification_pdf_path,
                product_summary=row.product_summary,
                remark=row.remark,
                item_type=row.item_type,
            )
            item_ids.append(item.id)

        self.upload_job_repo.update_job_status(job_id, "prepared", total_rows=len(rows))
        return item_ids
